from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATASETS_ROOT = PROJECT_ROOT / "datasets"
DEFAULT_OUTPUT = PROJECT_ROOT / "Health-Center-Data-Verification-Form.xlsx"

NAVY = "263238"
ROSE = "C95676"
PALE_ROSE = "FCEEF3"
PALE_BLUE = "EAF4FB"
PALE_GREEN = "EAF7EF"
PALE_YELLOW = "FFF7DB"
WHITE = "FFFFFF"
GRAY = "667085"
LIGHT_GRAY = "D0D5DD"
RED = "B42318"

SERVICES = [
    ("General Consultation", 20),
    ("Immunization", 15),
    ("Prenatal Care", 30),
    ("Blood Pressure Monitoring", 15),
    ("Child Wellness Check", 25),
    ("Family Planning", 25),
    ("Dental", 30),
]

SHOW_FOLDERS = [
    ("consultation", "General Consultation", "Each dated service actually rendered becomes one Show row."),
    ("family planning", "Family Planning", "Each dated family-planning service actually rendered becomes one Show row."),
    ("pediatric", "Per-visit review", "Vaccination becomes Immunization; general child assessment becomes Child Wellness Check."),
    ("prenatal", "Prenatal Care", "Each dated prenatal service actually rendered becomes one Show row."),
]

THIN = Side(style="thin", color=LIGHT_GRAY)
MEDIUM = Side(style="medium", color=NAVY)


def _image_inventory() -> tuple[list[tuple[str, int]], list[tuple[str, str, str]]]:
    show_counts = []
    for folder, _, _ in SHOW_FOLDERS:
        show_counts.append((folder, len(list((DATASETS_ROOT / folder).glob("*.jpg")))))

    no_show_dir = DATASETS_ROOT / "No-show"
    primary_files = sorted(
        (path for path in no_show_dir.glob("*.jpg") if path.stem.isdigit()),
        key=lambda path: int(path.stem),
    )
    supporting = {
        "1.jpg": "likod_sa_1.jpg",
        "2.jpg": "likod_sa_2.jpg",
    }
    no_show_rows = []
    for index, path in enumerate(primary_files, 1):
        no_show_rows.append((f"NS-{index:03d}", f"No-show/{path.name}", supporting.get(path.name, "")))
    return show_counts, no_show_rows


def _title(ws, text: str, subtitle: str, end_column: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    cell = ws.cell(1, 1, text)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(color=WHITE, bold=True, size=17)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    cell = ws.cell(2, 1, subtitle)
    cell.fill = PatternFill("solid", fgColor=PALE_ROSE)
    cell.font = Font(color=NAVY, italic=True, size=10)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 30


def _section(ws, row: int, text: str, end_column: int, color: str = PALE_BLUE) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_column)
    cell = ws.cell(row, 1, text)
    cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(color=NAVY, bold=True, size=11)
    cell.alignment = Alignment(vertical="center")
    cell.border = Border(bottom=MEDIUM)
    ws.row_dimensions[row].height = 24


def _headers(ws, row: int, values: list[str]) -> None:
    for column, value in enumerate(values, 1):
        cell = ws.cell(row, column, value)
        cell.fill = PatternFill("solid", fgColor=ROSE)
        cell.font = Font(color=WHITE, bold=True, size=9)
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[row].height = 38


def _boxed_row(ws, row: int, start_column: int, end_column: int, height: float = 25) -> None:
    for column in range(start_column, end_column + 1):
        cell = ws.cell(row, column)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = height


def _signature_line(ws, row: int, label: str, start_column: int, end_column: int) -> None:
    ws.merge_cells(start_row=row, start_column=start_column, end_row=row, end_column=end_column)
    cell = ws.cell(row, start_column, f"{label}: ______________________________________________")
    cell.font = Font(color=NAVY, size=10)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 27


def _print_setup(ws, orientation: str = "portrait", repeat_rows: str | None = None) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = orientation
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.45, bottom=0.45, header=0.2, footer=0.2)
    if repeat_rows:
        ws.print_title_rows = repeat_rows
    ws.oddFooter.center.text = "Health-center verification form | No patient identifiers permitted"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = GRAY
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.right.size = 8


def _add_table(ws, name: str, ref: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _guide_sheet(wb: Workbook, show_counts: list[tuple[str, int]], no_show_count: int) -> None:
    ws = wb.active
    ws.title = "Print Guide"
    _title(ws, "Health-Center Data Verification Pack", "Print this workbook and complete it with an authorized health-center representative. Do not write patient names or clinical details.", 7)

    _section(ws, 4, "What this pack verifies", 7)
    instructions = [
        "The exact study start and end dates.",
        "That every eligible dated rendered service in the approved period is a Show outcome.",
        "The exact missed appointment date and service for every clinic-confirmed No-Show source file.",
        "That Cancelled, rescheduled, blank, unreadable, or unverified outcomes are excluded.",
        "That only anonymized model fields may leave the local extraction environment or enter Google Colab.",
    ]
    for offset, text in enumerate(instructions, 5):
        ws.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=7)
        ws.cell(offset, 1, f"[ ] {text}")
        ws.cell(offset, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[offset].height = 25

    _section(ws, 11, "Current image inventory", 7, PALE_GREEN)
    _headers(ws, 12, ["Folder", "Class", "Image files", "How rows are created", "", "", ""])
    row = 13
    show_rules = {folder: rule for folder, _, rule in SHOW_FOLDERS}
    for folder, count in show_counts:
        ws.cell(row, 1, folder)
        ws.cell(row, 2, "Show")
        ws.cell(row, 3, count)
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=7)
        ws.cell(row, 4, show_rules[folder])
        _boxed_row(ws, row, 1, 7, 31)
        row += 1
    ws.cell(row, 1, "No-show")
    ws.cell(row, 2, "No-Show")
    ws.cell(row, 3, f"{no_show_count} primary + 2 back pages")
    ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=7)
    ws.cell(row, 4, "One row per exact missed appointment confirmed by the clinic; multiple rows are allowed for one source file.")
    _boxed_row(ws, row, 1, 7, 35)

    _section(ws, 20, "Completion order", 7, PALE_YELLOW)
    steps = [
        "1. Complete Study Approval and obtain the required signatures.",
        "2. Review and sign Show Rules.",
        "3. Complete every required cell in No-Show Verification; add continuation rows when one chart has multiple missed appointments.",
        "4. Review Field Checklist and Service Reference.",
        "5. Return the completed form for anonymized extraction and quality review.",
    ]
    for row, text in enumerate(steps, 21):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        ws.cell(row, 1, text)
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 25

    _section(ws, 27, "Privacy reminder", 7, PALE_ROSE)
    ws.merge_cells("A28:G30")
    ws["A28"] = "Do not add names, phone numbers, addresses, birth dates, diagnoses, clinical notes, signatures, Supabase IDs, or SMS details anywhere in this workbook. Source filenames are sufficient for verification. Raw images remain local and must never be uploaded to Google Colab."
    ws["A28"].font = Font(color=RED, bold=True, size=10)
    ws["A28"].alignment = Alignment(wrap_text=True, vertical="center")
    for column, width in enumerate([20, 15, 18, 23, 18, 18, 18], 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.print_area = "A1:G30"
    _print_setup(ws)


def _study_approval_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Study Approval")
    _title(ws, "Study Period and Data-Use Verification", "Complete all required fields before any image is converted into model-training rows.", 8)

    _section(ws, 4, "Study details", 8)
    fields = [
        (5, "Health center / unit", ""),
        (6, "Approved study start date (YYYY-MM-DD)", ""),
        (7, "Approved study end date (YYYY-MM-DD)", ""),
        (8, "Selection method", "All eligible records within the approved period"),
        (9, "No-Show evidence source", "Clinic-verified schedule/attendance records mapped by source filename"),
    ]
    for row, label, value in fields:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        ws.cell(row, 1, label)
        ws.cell(row, 1).font = Font(bold=True, color=NAVY)
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=8)
        ws.cell(row, 4, value)
        _boxed_row(ws, row, 1, 8, 29)

    _section(ws, 11, "Required confirmations", 8, PALE_YELLOW)
    confirmations = [
        "All eligible appointment records in the approved date range are represented, not a convenience selection.",
        "Every Show row requires a dated service actually rendered.",
        "Every No-Show row requires an exact scheduled date and clinic confirmation that the patient did not attend.",
        "Cancelled or rescheduled appointments are not No-Shows and will be excluded.",
        "Blank, ambiguous, unreadable, or unverified outcomes will be excluded rather than guessed.",
        "Historical visits outside the approved date range will not enter the study dataset.",
        "Identifiers may be viewed locally under authorization but will not enter the review workbook, CSV, Colab, model, or report.",
    ]
    for row, text in enumerate(confirmations, 12):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        ws.cell(row, 1, text)
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(row, 8, "[ ] Confirm")
        ws.cell(row, 8).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        _boxed_row(ws, row, 1, 8, 33)

    _section(ws, 20, "Authorization and signatures", 8, PALE_GREEN)
    _signature_line(ws, 21, "Authorized health-center verifier (printed name and role)", 1, 8)
    _signature_line(ws, 22, "Signature", 1, 5)
    _signature_line(ws, 22, "Date", 6, 8)
    _signature_line(ws, 24, "Researcher / data encoder", 1, 5)
    _signature_line(ws, 24, "Date", 6, 8)
    _signature_line(ws, 26, "Dean / adviser acknowledgment, if required", 1, 5)
    _signature_line(ws, 26, "Date", 6, 8)

    for column, width in enumerate([19, 19, 19, 18, 18, 18, 18, 15], 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.print_area = "A1:H27"
    _print_setup(ws)


def _show_rules_sheet(wb: Workbook, show_counts: list[tuple[str, int]]) -> None:
    ws = wb.create_sheet("Show Rules")
    _title(ws, "Show-Outcome Extraction Verification", "Confirm how dated rendered services will become appointment-level Show outcomes.", 7)
    _headers(ws, 4, ["Source folder", "Image count", "Canonical service", "Extraction rule", "Clinic confirms", "Verifier initials", "Date"])
    counts = dict(show_counts)
    row = 5
    for folder, service, rule in SHOW_FOLDERS:
        ws.cell(row, 1, folder)
        ws.cell(row, 2, counts[folder])
        ws.cell(row, 3, service)
        ws.cell(row, 4, rule)
        ws.cell(row, 5, "[ ] Yes  [ ] No")
        _boxed_row(ws, row, 1, 7, 50)
        row += 1

    _section(ws, 11, "Rules that apply to every Show folder", 7, PALE_YELLOW)
    rules = [
        "One dated rendered service equals one Show row; one image may produce multiple rows.",
        "Only visits inside the approved study period are included.",
        "The patient must have actually received the documented service; a registration date alone is not a Show.",
        "Pediatric rows are classified per visit as Immunization or Child Wellness Check based on the rendered service.",
        "Repeated visits for the same anonymous chart are retained in chronological order for prior-history calculation.",
        "Cancelled, planned-only, suggested return, and blank entries are excluded.",
    ]
    for row, text in enumerate(rules, 12):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        ws.cell(row, 1, text)
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(row, 6, "[ ] Confirm")
        ws.cell(row, 7, "Initials: ______")
        _boxed_row(ws, row, 1, 7, 32)

    _section(ws, 19, "Health-center verification", 7, PALE_GREEN)
    _signature_line(ws, 20, "Verifier printed name and role", 1, 7)
    _signature_line(ws, 21, "Signature", 1, 4)
    _signature_line(ws, 21, "Date", 5, 7)
    for column, width in enumerate([20, 13, 27, 55, 18, 17, 15], 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.print_area = "A1:G22"
    _print_setup(ws, "landscape")


def _no_show_sheet(wb: Workbook, no_show_rows: list[tuple[str, str, str]]) -> None:
    ws = wb.create_sheet("No-Show Verification")
    _title(ws, "Clinic-Confirmed No-Show Verification", "Complete every REQUIRED field. Use the blank continuation rows when one source chart has more than one confirmed missed appointment.", 12)
    headers = [
        "Record ref",
        "Primary source file",
        "Supporting back page",
        "Missed appointment date REQUIRED",
        "Service REQUIRED",
        "Scheduled? REQUIRED",
        "Did not attend? REQUIRED",
        "Cancelled or rescheduled? REQUIRED",
        "Scheduled time OPTIONAL",
        "Booking date OPTIONAL",
        "Verifier initials REQUIRED",
        "Verification date REQUIRED",
    ]
    _headers(ws, 4, headers)
    row = 5
    for record_ref, source, supporting in no_show_rows:
        ws.cell(row, 1, record_ref)
        ws.cell(row, 2, source)
        ws.cell(row, 3, supporting)
        _boxed_row(ws, row, 1, 12, 42)
        row += 1
    for _ in range(5):
        ws.cell(row, 1, "CONT")
        _boxed_row(ws, row, 1, 12, 42)
        row += 1
    last_row = row - 1

    date_validation = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2035,12,31)", allow_blank=True)
    service_validation = DataValidation(type="list", formula1="=ServiceNames", allow_blank=True)
    yes_no_validation = DataValidation(type="list", formula1="=YesNoValues", allow_blank=True)
    time_validation = DataValidation(type="time", operator="between", formula1="TIME(0,0,0)", formula2="TIME(23,59,59)", allow_blank=True)
    for validation in [date_validation, service_validation, yes_no_validation, time_validation]:
        ws.add_data_validation(validation)
    date_validation.add(f"D5:D{last_row}")
    date_validation.add(f"J5:J{last_row}")
    date_validation.add(f"L5:L{last_row}")
    service_validation.add(f"E5:E{last_row}")
    yes_no_validation.add(f"F5:H{last_row}")
    time_validation.add(f"I5:I{last_row}")

    for current_row in range(5, last_row + 1):
        ws.cell(current_row, 4).number_format = "yyyy-mm-dd"
        ws.cell(current_row, 9).number_format = "hh:mm"
        ws.cell(current_row, 10).number_format = "yyyy-mm-dd"
        ws.cell(current_row, 12).number_format = "yyyy-mm-dd"

    _add_table(ws, "NoShowVerificationTable", f"A4:L{last_row}")
    note_row = last_row + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 2, end_column=12)
    ws.cell(note_row, 1, "VALID NO-SHOW RULE: Scheduled = Yes, Did not attend = Yes, Cancelled/rescheduled = No. The missed date and service must refer to the same appointment. Do not infer a missed date from a blank chart row, suggested return date, photo date, or latest recorded visit.")
    ws.cell(note_row, 1).fill = PatternFill("solid", fgColor=PALE_YELLOW)
    ws.cell(note_row, 1).font = Font(color=RED, bold=True, size=10)
    ws.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.print_area = f"A1:L{note_row + 2}"

    widths = [13, 22, 22, 19, 27, 14, 17, 20, 16, 16, 17, 17]
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "A5"
    _print_setup(ws, "landscape", "$1:$4")


def _field_checklist_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Field Checklist")
    _title(ws, "Extraction Field Checklist", "This defines what is needed, what is optional, and what must never enter the dataset.", 6)
    headers = ["Field", "Requirement", "Source", "Final ML CSV", "Verification rule", "Health-center notes (no patient details)"]
    _headers(ws, 4, headers)
    rows = [
        ("appointment_date", "Required", "Dated rendered service or verified missed schedule", "Retained", "Exact date inside approved study period"),
        ("age", "Required", "Age on form or locally calculated from birth date", "Retained", "Age at appointment, 0-130; birth date removed"),
        ("gender", "Required", "Patient chart", "Retained", "Female or Male under current system schema"),
        ("service_name", "Required", "Rendered service or verified missed schedule", "Retained", "Use one canonical service"),
        ("duration_minutes", "Derived", "Canonical service reference", "Retained", "Expected duration, not observed consultation length"),
        ("attendance outcome", "Required", "Rendered service or clinic-confirmed No-Show", "Converted to label", "Show = 0; No-Show = 1"),
        ("anonymous chart code", "Temporary", "Assigned locally per chart", "Removed", "Used only to calculate strictly earlier history"),
        ("prior_appointments", "Derived", "Earlier verified outcomes", "Retained", "Count only outcomes before current appointment"),
        ("prior_no_show_rate", "Derived", "Earlier verified outcomes", "Retained", "Earlier No-Shows / earlier outcomes; 0 if none"),
        ("day_of_week", "Derived", "Appointment date", "Retained", "Calculated automatically"),
        ("is_weekend", "Derived", "Appointment date", "Retained", "Calculated automatically"),
        ("scheduled time", "Optional", "Clinic schedule", "Excluded from initial real model", "Collect only if available consistently"),
        ("booking date", "Optional", "Clinic schedule", "Excluded from initial real model", "Collect only if available consistently"),
        ("names / contacts / address", "Prohibited", "Patient chart", "Never retained", "Local viewing only under authorization"),
        ("diagnoses / notes / signatures", "Prohibited", "Patient chart", "Never retained", "Do not transcribe"),
        ("SMS / prediction outcomes", "Prohibited", "System records", "Never retained", "Post-appointment leakage"),
    ]
    for row_number, values in enumerate(rows, 5):
        for column, value in enumerate(values, 1):
            ws.cell(row_number, column, value)
        _boxed_row(ws, row_number, 1, 6, 39)
        if values[1] == "Prohibited":
            for column in range(1, 7):
                ws.cell(row_number, column).fill = PatternFill("solid", fgColor=PALE_ROSE)
                ws.cell(row_number, column).font = Font(color=RED, bold=column == 1)
    _add_table(ws, "FieldChecklistTable", f"A4:F{len(rows) + 4}")
    for column, width in enumerate([27, 16, 37, 25, 47, 32], 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.print_area = f"A1:F{len(rows) + 4}"
    _print_setup(ws, "landscape", "$1:$4")


def _service_reference_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Service Reference")
    _title(ws, "Canonical Health-Service Reference", "Use these exact service names and expected durations in the verification and final dataset.", 5)
    _headers(ws, 4, ["Canonical service", "Expected duration (minutes)", "Common source / alias", "Use for these records", "Health-center confirmation"])
    aliases = {
        "General Consultation": "consultation",
        "Immunization": "pediatric vaccination; Pediatric / Immunization",
        "Prenatal Care": "prenatal; Maternal Care - Prenatal",
        "Blood Pressure Monitoring": "blood-pressure monitoring",
        "Child Wellness Check": "pediatric general child assessment",
        "Family Planning": "family planning",
        "Dental": "Dental Checkup",
    }
    for row, (service, duration) in enumerate(SERVICES, 5):
        ws.cell(row, 1, service)
        ws.cell(row, 2, duration)
        ws.cell(row, 3, aliases[service])
        ws.cell(row, 4, "Use only when this exact service was rendered or scheduled.")
        ws.cell(row, 5, "[ ] Confirm")
        _boxed_row(ws, row, 1, 5, 37)
    _add_table(ws, "ServiceReferenceTable", "A4:E11")
    for column, width in enumerate([31, 23, 42, 54, 23], 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.print_area = "A1:E12"
    _print_setup(ws, "landscape")


def _lists_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Lists")
    ws.append(["ServiceNames", "YesNoValues"])
    for index, (service, _) in enumerate(SERVICES, 2):
        ws.cell(index, 1, service)
    ws.cell(2, 2, "Yes")
    ws.cell(3, 2, "No")
    wb.defined_names.add(DefinedName("ServiceNames", attr_text="Lists!$A$2:$A$8"))
    wb.defined_names.add(DefinedName("YesNoValues", attr_text="Lists!$B$2:$B$3"))
    ws.sheet_state = "hidden"


def create_workbook(output_path: Path = DEFAULT_OUTPUT) -> Path:
    show_counts, no_show_rows = _image_inventory()
    if len(no_show_rows) != 15:
        raise RuntimeError(f"Expected 15 primary No-Show files, found {len(no_show_rows)}.")

    wb = Workbook()
    wb.properties.title = "Health-Center Data Verification Form"
    wb.properties.subject = "Appointment-level Show and No-Show dataset verification"
    wb.properties.creator = "Smart Queue capstone project"
    _guide_sheet(wb, show_counts, len(no_show_rows))
    _study_approval_sheet(wb)
    _show_rules_sheet(wb, show_counts)
    _no_show_sheet(wb, no_show_rows)
    _field_checklist_sheet(wb)
    _service_reference_sheet(wb)
    _lists_sheet(wb)
    wb.save(output_path)

    verified = load_workbook(output_path, data_only=False)
    expected = [
        "Print Guide",
        "Study Approval",
        "Show Rules",
        "No-Show Verification",
        "Field Checklist",
        "Service Reference",
        "Lists",
    ]
    if verified.sheetnames != expected:
        raise RuntimeError(f"Unexpected workbook sheets: {verified.sheetnames}")
    verified.close()
    return output_path


if __name__ == "__main__":
    print(create_workbook())
