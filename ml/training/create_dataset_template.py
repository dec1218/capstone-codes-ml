from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT = PROJECT_ROOT / "No-Show-Dataset-Template.xlsx"

NAVY = "263238"
ROSE = "C95676"
PALE_ROSE = "FCEEF3"
PALE_BLUE = "EAF4FB"
PALE_GREEN = "EAF7EF"
PALE_YELLOW = "FFF7DB"
WHITE = "FFFFFF"
GRAY = "667085"
LIGHT_GRAY = "E4E7EC"
RED = "B42318"
GREEN = "067647"

SERVICES = {
    "General Consultation": 20,
    "Immunization": 15,
    "Prenatal Care": 30,
    "Blood Pressure Monitoring": 15,
    "Child Wellness Check": 25,
    "Family Planning": 25,
    "Dental": 30,
}
GENDERS = ["Female", "Male"]
ATTENDANCE_STATUSES = ["Present", "Late", "No-Show"]
VERIFICATION_VALUES = ["Yes", "No"]

RAW_HEADERS = [
    "record_code",
    "anonymous_patient_code",
    "appointment_date",
    "appointment_time",
    "created_at",
    "age",
    "gender",
    "service_name",
    "duration_minutes",
    "attendance_status",
    "data_verified",
]

ML_HEADERS = [
    "appointment_date",
    "appointment_time",
    "created_at",
    "age",
    "gender",
    "service_name",
    "duration_minutes",
    "prior_appointments",
    "prior_no_show_rate",
    "label",
]

HEADER_COMMENTS = {
    "record_code": "A non-identifying row reference such as REC-0001. Do not use a live appointment UUID.",
    "anonymous_patient_code": "A stable anonymous code used only to calculate earlier attendance. Remove it from the ML-ready export.",
    "appointment_date": "The scheduled appointment date, not the attendance-recording date.",
    "appointment_time": "The scheduled appointment time in local health-center time.",
    "created_at": "When the appointment was booked. This must be before or equal to the appointment date.",
    "age": "Age known before the appointment. Allowed range: 0 to 130.",
    "gender": "Recorded category supported by the current system: Female or Male.",
    "service_name": "Use one canonical service from the dropdown. Unknown labels require review.",
    "duration_minutes": "Expected service duration. Use the canonical value from Valid Values.",
    "attendance_status": "Verified actual outcome: Present, Late, or No-Show. Never infer No-Show from a blank value.",
    "data_verified": "Yes only when the appointment and attendance outcome were checked against an authorized source.",
    "prior_appointments": "Number of this anonymous patient's verified outcomes strictly before this appointment.",
    "prior_no_show_rate": "Earlier No-Shows divided by earlier verified outcomes. Use 0 when there are no earlier outcomes.",
    "label": "Binary target: 1 = No-Show; 0 = Present or Late.",
}


def _example_rows() -> list[dict[str, object]]:
    raw = [
        ("REC-0001", "ANON-001", date(2026, 1, 5), time(8, 0), datetime(2025, 12, 20, 9, 15), 34, "Female", "General Consultation", "Present"),
        ("REC-0002", "ANON-002", date(2026, 1, 6), time(9, 30), datetime(2025, 12, 28, 14, 10), 51, "Male", "Blood Pressure Monitoring", "Late"),
        ("REC-0003", "ANON-003", date(2026, 1, 8), time(10, 0), datetime(2026, 1, 2, 11, 30), 6, "Female", "Immunization", "Present"),
        ("REC-0004", "ANON-004", date(2026, 1, 10), time(13, 0), datetime(2025, 12, 15, 8, 45), 28, "Female", "Prenatal Care", "No-Show"),
        ("REC-0005", "ANON-005", date(2026, 1, 12), time(14, 30), datetime(2026, 1, 7, 10, 0), 42, "Male", "Dental", "Present"),
        ("REC-0006", "ANON-006", date(2026, 1, 14), time(11, 0), datetime(2025, 12, 30, 15, 20), 31, "Female", "Family Planning", "Present"),
        ("REC-0007", "ANON-007", date(2026, 1, 16), time(8, 30), datetime(2026, 1, 3, 9, 5), 3, "Male", "Child Wellness Check", "No-Show"),
        ("REC-0008", "ANON-008", date(2026, 1, 19), time(15, 0), datetime(2026, 1, 10, 13, 40), 63, "Female", "General Consultation", "Present"),
        ("REC-0009", "ANON-001", date(2026, 2, 10), time(9, 0), datetime(2026, 1, 18, 8, 20), 34, "Female", "Dental", "No-Show"),
        ("REC-0010", "ANON-002", date(2026, 2, 15), time(10, 30), datetime(2026, 2, 1, 16, 10), 51, "Male", "General Consultation", "Present"),
        ("REC-0011", "ANON-003", date(2026, 2, 18), time(13, 30), datetime(2026, 2, 5, 9, 0), 6, "Female", "Child Wellness Check", "Late"),
        ("REC-0012", "ANON-004", date(2026, 2, 20), time(8, 0), datetime(2026, 2, 2, 10, 45), 28, "Female", "Prenatal Care", "Present"),
        ("REC-0013", "ANON-005", date(2026, 2, 23), time(14, 0), datetime(2026, 2, 19, 11, 50), 42, "Male", "Dental", "Present"),
        ("REC-0014", "ANON-006", date(2026, 2, 25), time(9, 30), datetime(2026, 2, 8, 14, 15), 31, "Female", "Family Planning", "No-Show"),
        ("REC-0015", "ANON-007", date(2026, 3, 2), time(11, 30), datetime(2026, 2, 20, 8, 30), 3, "Male", "Immunization", "Present"),
        ("REC-0016", "ANON-008", date(2026, 3, 5), time(15, 30), datetime(2026, 2, 15, 13, 0), 63, "Female", "Blood Pressure Monitoring", "Late"),
        ("REC-0017", "ANON-001", date(2026, 3, 20), time(8, 30), datetime(2026, 3, 1, 9, 40), 34, "Female", "General Consultation", "Present"),
        ("REC-0018", "ANON-002", date(2026, 4, 1), time(10, 0), datetime(2026, 3, 10, 15, 25), 51, "Male", "Dental", "No-Show"),
        ("REC-0019", "ANON-003", date(2026, 4, 4), time(13, 0), datetime(2026, 3, 25, 10, 10), 6, "Female", "Immunization", "Present"),
        ("REC-0020", "ANON-004", date(2026, 4, 8), time(9, 0), datetime(2026, 3, 15, 8, 55), 28, "Female", "Prenatal Care", "Present"),
        ("REC-0021", "ANON-005", date(2026, 4, 11), time(14, 30), datetime(2026, 4, 2, 11, 20), 42, "Male", "Dental", "Late"),
        ("REC-0022", "ANON-006", date(2026, 4, 14), time(8, 0), datetime(2026, 3, 30, 16, 0), 31, "Female", "Family Planning", "Present"),
        ("REC-0023", "ANON-007", date(2026, 4, 18), time(10, 30), datetime(2026, 4, 1, 9, 35), 3, "Male", "Child Wellness Check", "No-Show"),
        ("REC-0024", "ANON-008", date(2026, 4, 22), time(15, 0), datetime(2026, 4, 12, 12, 5), 63, "Female", "General Consultation", "Present"),
    ]
    return [{
        "record_code": record_code,
        "anonymous_patient_code": patient_code,
        "appointment_date": appointment_date,
        "appointment_time": appointment_time,
        "created_at": created_at,
        "age": age,
        "gender": gender,
        "service_name": service,
        "duration_minutes": SERVICES[service],
        "attendance_status": attendance,
        "data_verified": "Yes",
    } for record_code, patient_code, appointment_date, appointment_time, created_at, age, gender, service, attendance in raw]


def _ml_ready_rows(raw_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    history: dict[str, list[int]] = defaultdict(list)
    output = []
    for row in sorted(raw_rows, key=lambda item: (item["appointment_date"], item["appointment_time"])):
        patient_history = history[str(row["anonymous_patient_code"])]
        no_show_count = sum(patient_history)
        output.append({
            "appointment_date": row["appointment_date"],
            "appointment_time": row["appointment_time"],
            "created_at": row["created_at"],
            "age": row["age"],
            "gender": row["gender"],
            "service_name": row["service_name"],
            "duration_minutes": row["duration_minutes"],
            "prior_appointments": len(patient_history),
            "prior_no_show_rate": no_show_count / len(patient_history) if patient_history else 0,
            "label": int(row["attendance_status"] == "No-Show"),
        })
        patient_history.append(int(row["attendance_status"] == "No-Show"))
    return output


def _style_title(ws, title: str, subtitle: str, end_column: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    title_cell = ws.cell(1, 1, title)
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.font = Font(color=WHITE, bold=True, size=18)
    title_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    subtitle_cell = ws.cell(2, 1, subtitle)
    subtitle_cell.fill = PatternFill("solid", fgColor=PALE_ROSE)
    subtitle_cell.font = Font(color=NAVY, italic=True, size=10)
    subtitle_cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 31


def _style_header_row(ws, row_number: int, headers: list[str]) -> None:
    thin = Side(style="thin", color=LIGHT_GRAY)
    for column, header in enumerate(headers, 1):
        cell = ws.cell(row_number, column, header)
        cell.fill = PatternFill("solid", fgColor=ROSE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=thin)
        if header in HEADER_COMMENTS:
            cell.comment = Comment(HEADER_COMMENTS[header], "Capstone template")
    ws.row_dimensions[row_number].height = 34


def _format_data_sheet(ws, header_row: int, last_row: int, widths: list[int]) -> None:
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{ws.cell(header_row, len(widths)).coordinate}"
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in range(header_row + 1, last_row + 1):
        ws.cell(row, 3).number_format = "yyyy-mm-dd"
        ws.cell(row, 4).number_format = "hh:mm"
        ws.cell(row, 5).number_format = "yyyy-mm-dd hh:mm"
        for column in range(1, len(widths) + 1):
            ws.cell(row, column).alignment = Alignment(vertical="top")


def _add_table(ws, name: str, ref: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)


def _instructions_sheet(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Start Here"
    _style_title(ws, "Real No-Show Dataset Guide", "Collection template and fictional examples for the barangay health-center model. This workbook is not a training dataset.", 6)
    sections = [
        (4, "Purpose", PALE_BLUE, [
            "Use Blank Collection to understand the fields needed for verified real appointment outcomes.",
            "Fictional Example demonstrates raw rows; ML-Ready Example demonstrates the de-identified Colab input.",
            "Keep the original source unchanged. Create a reviewed anonymized export for model training.",
        ]),
        (10, "Privacy rules", PALE_YELLOW, [
            "Never include patient names, phone numbers, addresses, email, diagnoses, free-text clinical notes, Supabase UUIDs, or SMS results.",
            "Use a temporary anonymous patient code only to calculate prior history; remove it from the final ML-ready CSV.",
            "Never upload ml/.env, Supabase service-role keys, PhilSMS tokens, or other credentials to Google Colab.",
        ]),
        (16, "Outcome rules", PALE_GREEN, [
            "Present and Late become label 0 (Show). Verified No-Show becomes label 1.",
            "Do not convert blanks, missing attendance, Scheduled appointments, or Cancelled appointments into No-Show.",
            "Only use records whose data_verified value is Yes.",
        ]),
        (22, "Leakage rules", PALE_ROSE, [
            "prior_appointments and prior_no_show_rate must use outcomes strictly before the current appointment.",
            "Do not use check-in/out times, attendance_status, reminder delivery, probability, risk level, or prediction result as model features.",
            "Split data chronologically: oldest 70% train, next 15% validation, newest 15% test.",
        ]),
        (28, "Workflow", PALE_BLUE, [
            "1. Collect and verify one row per completed appointment.",
            "2. Review duplicates, missing values, services, dates, and both outcome classes.",
            "3. Calculate prior history using anonymous codes, then remove all record/patient codes.",
            "4. Export only the ML-ready columns to CSV and train in Google Colab.",
            "5. Keep the current SMS reminder policy independent until a real model is validated and approved.",
        ]),
    ]
    for start_row, heading, fill_color, lines in sections:
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=6)
        heading_cell = ws.cell(start_row, 1, heading)
        heading_cell.fill = PatternFill("solid", fgColor=fill_color)
        heading_cell.font = Font(color=NAVY, bold=True, size=12)
        for offset, line in enumerate(lines, 1):
            ws.merge_cells(start_row=start_row + offset, start_column=1, end_row=start_row + offset, end_column=6)
            cell = ws.cell(start_row + offset, 1, f"- {line}")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(color=NAVY, size=10)
            ws.row_dimensions[start_row + offset].height = 27
    ws.column_dimensions["A"].width = 24
    for column in "BCDEF":
        ws.column_dimensions[column].width = 18
    ws.sheet_view.showGridLines = False


def _valid_values_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Valid Values")
    columns = {
        "Service": list(SERVICES),
        "Duration": list(SERVICES.values()),
        "Gender": GENDERS,
        "Attendance": ATTENDANCE_STATUSES,
        "Verified": VERIFICATION_VALUES,
    }
    for column, (heading, values) in enumerate(columns.items(), 1):
        ws.cell(1, column, heading)
        ws.cell(1, column).fill = PatternFill("solid", fgColor=NAVY)
        ws.cell(1, column).font = Font(color=WHITE, bold=True)
        for row, value in enumerate(values, 2):
            ws.cell(row, column, value)
        ws.column_dimensions[ws.cell(1, column).column_letter].width = 30 if column == 1 else 18
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    wb.defined_names.add(DefinedName("ServiceNames", attr_text="'Valid Values'!$A$2:$A$8"))
    wb.defined_names.add(DefinedName("GenderValues", attr_text="'Valid Values'!$C$2:$C$3"))
    wb.defined_names.add(DefinedName("AttendanceValues", attr_text="'Valid Values'!$D$2:$D$4"))
    wb.defined_names.add(DefinedName("VerificationValues", attr_text="'Valid Values'!$E$2:$E$3"))


def _blank_template_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Blank Collection")
    _style_title(ws, "Blank Real-Data Collection Template", "Enter verified, anonymized appointment outcomes only. Do not use the fictional examples as real data.", len(RAW_HEADERS))
    _style_header_row(ws, 4, RAW_HEADERS)
    for row in range(5, 55):
        for column in range(1, len(RAW_HEADERS) + 1):
            ws.cell(row, column, None)
    _format_data_sheet(ws, 4, 54, [15, 23, 18, 18, 22, 10, 14, 31, 20, 21, 17])
    _add_table(ws, "BlankCollectionTable", f"A4:K54")

    date_validation = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2035,12,31)", allow_blank=False)
    time_validation = DataValidation(type="time", operator="between", formula1="TIME(7,0,0)", formula2="TIME(17,0,0)", allow_blank=False)
    age_validation = DataValidation(type="whole", operator="between", formula1="0", formula2="130", allow_blank=False)
    duration_validation = DataValidation(type="whole", operator="between", formula1="1", formula2="240", allow_blank=False)
    gender_validation = DataValidation(type="list", formula1="=GenderValues", allow_blank=False)
    service_validation = DataValidation(type="list", formula1="=ServiceNames", allow_blank=False)
    attendance_validation = DataValidation(type="list", formula1="=AttendanceValues", allow_blank=False)
    verified_validation = DataValidation(type="list", formula1="=VerificationValues", allow_blank=False)
    for validation, cell_range in [
        (date_validation, "C5:C54"),
        (time_validation, "D5:D54"),
        (age_validation, "F5:F54"),
        (gender_validation, "G5:G54"),
        (service_validation, "H5:H54"),
        (duration_validation, "I5:I54"),
        (attendance_validation, "J5:J54"),
        (verified_validation, "K5:K54"),
    ]:
        validation.error = "Use the required format or select an approved value."
        validation.errorTitle = "Invalid value"
        validation.prompt = "See Start Here and Data Dictionary before entering data."
        validation.promptTitle = "Verified anonymized data only"
        validation.showErrorMessage = True
        validation.showInputMessage = True
        ws.add_data_validation(validation)
        validation.add(cell_range)
    ws.sheet_view.showGridLines = False


def _example_sheet(wb: Workbook, rows: list[dict[str, object]]) -> None:
    ws = wb.create_sheet("Fictional Example")
    _style_title(ws, "Fictional Raw Collection Example", "All rows and anonymous codes are invented for format demonstration. Do not combine this sheet with real training data.", len(RAW_HEADERS))
    _style_header_row(ws, 4, RAW_HEADERS)
    for row_number, row in enumerate(rows, 5):
        for column, header in enumerate(RAW_HEADERS, 1):
            ws.cell(row_number, column, row[header])
    last_row = len(rows) + 4
    _format_data_sheet(ws, 4, last_row, [15, 23, 18, 18, 22, 10, 14, 31, 20, 21, 17])
    _add_table(ws, "FictionalCollectionTable", f"A4:K{last_row}")
    ws.conditional_formatting.add(f"J5:J{last_row}", CellIsRule(operator="equal", formula=['"No-Show"'], fill=PatternFill("solid", fgColor=PALE_ROSE), font=Font(color=RED, bold=True)))
    ws.sheet_view.showGridLines = False


def _ml_ready_sheet(wb: Workbook, rows: list[dict[str, object]]) -> None:
    ws = wb.create_sheet("ML-Ready Example")
    _style_title(ws, "Fictional ML-Ready Example", "This is the exact column shape expected by Colab. Record and patient codes have been removed after prior-history calculation.", len(ML_HEADERS))
    _style_header_row(ws, 4, ML_HEADERS)
    for row_number, row in enumerate(rows, 5):
        for column, header in enumerate(ML_HEADERS, 1):
            ws.cell(row_number, column, row[header])
        ws.cell(row_number, 9).number_format = "0.000"
    last_row = len(rows) + 4
    _format_data_sheet(ws, 4, last_row, [18, 18, 22, 10, 14, 31, 20, 21, 22, 12])
    _add_table(ws, "MLReadyExampleTable", f"A4:J{last_row}")
    ws.conditional_formatting.add(f"J5:J{last_row}", CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", fgColor=PALE_ROSE), font=Font(color=RED, bold=True)))
    ws.conditional_formatting.add(f"J5:J{last_row}", CellIsRule(operator="equal", formula=["0"], fill=PatternFill("solid", fgColor=PALE_GREEN), font=Font(color=GREEN, bold=True)))
    ws.sheet_view.showGridLines = False


def _dictionary_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Data Dictionary")
    headers = ["Field", "Raw collection", "ML-ready CSV", "Type / format", "Allowed values or example", "Purpose and rule"]
    _style_title(ws, "Data Dictionary", "Use these definitions consistently across paper transcription, Supabase export, and Google Colab.", len(headers))
    _style_header_row(ws, 4, headers)
    rows = [
        ("record_code", "Required", "Remove", "Text", "REC-0001", "Non-identifying audit reference; never use a live appointment UUID."),
        ("anonymous_patient_code", "Required for history", "Remove", "Text", "ANON-001", "Stable anonymous grouping key used only to compute earlier attendance."),
        ("appointment_date", "Required", "Required", "Date", "2026-01-05", "Scheduled service date."),
        ("appointment_time", "Required", "Required", "Time", "08:00", "Scheduled local time, normally 07:00-17:00."),
        ("created_at", "Required", "Required", "Date/time", "2025-12-20 09:15", "Booking time used to calculate lead days; must not occur after the appointment."),
        ("age", "Required", "Required", "Whole number", "0-130", "Age available before the appointment."),
        ("gender", "Required", "Required", "Category", "Female, Male", "Recorded category supported by the current system."),
        ("service_name", "Required", "Required", "Category", "Canonical service dropdown", "Unknown or blank services must be reviewed, not guessed."),
        ("duration_minutes", "Required", "Required", "Whole number", "15, 20, 25, 30", "Expected service duration, not observed consultation length."),
        ("attendance_status", "Required", "Convert then remove", "Category", "Present, Late, No-Show", "Verified outcome. Present/Late -> 0; No-Show -> 1."),
        ("data_verified", "Required", "Remove", "Yes/No", "Yes", "Only Yes rows may enter the official dataset."),
        ("prior_appointments", "Derived", "Required", "Whole number", "0 or greater", "Count of verified outcomes strictly before this appointment."),
        ("prior_no_show_rate", "Derived", "Required", "Decimal", "0.000-1.000", "Earlier No-Shows divided by earlier verified outcomes."),
        ("label", "Derived", "Required", "Binary", "0 or 1", "Target only: 0 = Present/Late; 1 = No-Show."),
    ]
    for row_number, values in enumerate(rows, 5):
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_number, column, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_number].height = 34
    _add_table(ws, "DataDictionaryTable", f"A4:F{len(rows) + 4}")
    widths = [26, 21, 21, 19, 29, 62]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False


def create_workbook(output_path: Path = DEFAULT_OUTPUT) -> Path:
    wb = Workbook()
    _instructions_sheet(wb)
    _valid_values_sheet(wb)
    _blank_template_sheet(wb)
    examples = _example_rows()
    _example_sheet(wb, examples)
    _ml_ready_sheet(wb, _ml_ready_rows(examples))
    _dictionary_sheet(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(output_path)

    # Reopen before returning so generation fails immediately if the XLSX package is invalid.
    verified = load_workbook(output_path, data_only=False)
    expected = ["Start Here", "Valid Values", "Blank Collection", "Fictional Example", "ML-Ready Example", "Data Dictionary"]
    if verified.sheetnames != expected:
        raise RuntimeError(f"Unexpected workbook sheets: {verified.sheetnames}")
    verified.close()
    return output_path


if __name__ == "__main__":
    print(create_workbook())
